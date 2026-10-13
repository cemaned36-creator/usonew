
from fastapi import APIRouter, Depends, HTTPException
from sqlalchemy.orm import Session

from app.database import get_db
from app.models.lesson import Lesson
from app.models.room import Room
from app.models.weekly_schedule_entry import WeeklyScheduleEntry
from app.schemas.weekly_schedule import WeeklyScheduleCreate
from app.services.access import calculate_transition_time

from io import BytesIO

from fastapi import UploadFile, File
from openpyxl import load_workbook

from app.models.subject import Subject
from app.models.teacher import Teacher
from app.models.student_group import StudentGroup

from datetime import time, datetime

router = APIRouter(
    prefix="/weekly-schedule",
    tags=["Weekly Schedule"],
)


def parse_time_value(value):
    """Преобразует время из Excel в datetime.time."""
    if isinstance(value, time):
        return value

    if isinstance(value, datetime):
        return value.time()

    if isinstance(value, str):
        value = value.strip()

        for fmt in ("%H:%M", "%H:%M:%S"):
            try:
                return datetime.strptime(value, fmt).time()
            except ValueError:
                continue

    return None


def parse_weekday(value):
    """Возвращает номер дня недели: 0 — понедельник, 6 — воскресенье."""
    if isinstance(value, bool) or value is None:
        return None

    if isinstance(value, int) and 0 <= value <= 6:
        return value

    weekdays = {
        "понедельник": 0,
        "вторник": 1,
        "среда": 2,
        "четверг": 3,
        "пятница": 4,
        "суббота": 5,
        "воскресенье": 6,
    }

    text = str(value).strip().casefold()

    if text in weekdays:
        return weekdays[text]

    try:
        number = int(text)
        if 0 <= number <= 6:
            return number
    except ValueError:
        pass

    return None


@router.post("/import/preview")
async def preview_schedule_import(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    # 1. Проверяем расширение файла
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(
            status_code=400,
            detail="Загрузите файл Excel с расширением .xlsx",
        )

    # 2. Читаем Excel
    content = await file.read()

    try:
        workbook = load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=True,
        )
    except Exception:
        raise HTTPException(
            status_code=400,
            detail="Не удалось прочитать Excel-файл",
        )

    # 3. Проверяем нужный лист
    sheet_name = "Данные для импорта"

    if sheet_name not in workbook.sheetnames:
        workbook.close()
        raise HTTPException(
            status_code=400,
            detail=f"В файле отсутствует лист «{sheet_name}»",
        )

    sheet = workbook[sheet_name]
    rows = sheet.iter_rows(values_only=True)

    headers = next(rows, None)

    if not headers:
        workbook.close()
        raise HTTPException(
            status_code=400,
            detail="Лист Excel пуст",
        )

    # Приводим названия столбцов к единому виду
    headers = [
        str(value).strip().lower() if value is not None else ""
        for value in headers
    ]

    required = {
        "weekday",
        "group",
        "start_time",
        "end_time",
        "subject",
        "teacher",
        "room",
    }

    missing = required - set(headers)

    if missing:
        workbook.close()
        raise HTTPException(
            status_code=400,
            detail={
                "message": "В Excel отсутствуют обязательные столбцы",
                "missing_columns": sorted(missing),
            },
        )

    columns = {name: headers.index(name) for name in required}

    def get_cell(row, column_name):
        index = columns[column_name]
        return row[index] if index < len(row) else None

    # 4. Загружаем справочники из базы
    groups = {
        str(item.name).strip().casefold(): item
        for item in db.query(StudentGroup).all()
    }

    subjects = {
        str(item.name).strip().casefold(): item
        for item in db.query(Subject).all()
    }

    teachers = {
        str(item.name).strip().casefold(): item
        for item in db.query(Teacher).all()
    }

    rooms = {
        str(item.name).strip().casefold(): item
        for item in db.query(Room).all()
    }

    result = []
    row_number = 1  # строка заголовков

    # 5. Проверяем строки Excel
    for row in rows:
        row_number += 1

        if not row or all(value is None for value in row):
            continue

        group_name = get_cell(row, "group")
        subject_name = get_cell(row, "subject")
        teacher_name = get_cell(row, "teacher")
        room_name = get_cell(row, "room")
        start_time_raw = get_cell(row, "start_time")
        end_time_raw = get_cell(row, "end_time")
        weekday_raw = get_cell(row, "weekday")

        start_time = parse_time_value(start_time_raw)
        end_time = parse_time_value(end_time_raw)
        weekday = parse_weekday(weekday_raw)

        issues = []

        def find_item(value, directory, label):
            if value is None or not str(value).strip():
                issues.append(f"Не указано: {label}")
                return None

            key = str(value).strip().casefold()
            item = directory.get(key)

            if item is None:
                issues.append(
                    f"{label} «{str(value).strip()}» не найден в базе"
                )

            return item

        group = find_item(group_name, groups, "Группа")
        subject = find_item(subject_name, subjects, "Предмет")
        teacher = find_item(teacher_name, teachers, "Преподаватель")
        room = find_item(room_name, rooms, "Аудитория")

        if start_time is None:
            issues.append("Время начала отсутствует или имеет неверный формат")

        if end_time is None:
            issues.append("Время окончания отсутствует или имеет неверный формат")

        if start_time is not None and end_time is not None:
            if start_time >= end_time:
                issues.append("Время окончания должно быть позже начала")

        if weekday is None:
            issues.append(
                "День недели должен быть числом от 0 до 6 "
                "или названием дня недели на русском языке"
            )

        # Ищем занятие, соответствующее группе,
        # предмету и преподавателю
        lesson = None

        if group and subject and teacher:
            lesson = (
                db.query(Lesson)
                .filter(
                    Lesson.group_id == group.id,
                    Lesson.subject_id == subject.id,
                    Lesson.teacher_id == teacher.id,
                )
                .first()
            )

            if lesson is None:
                issues.append(
                    "В базе нет занятия для этой группы, "
                    "предмета и преподавателя"
                )

        result.append({
            "excel_row": row_number,
            "weekday": str(weekday) if weekday is not None else None,
            "group": str(group_name) if group_name is not None else None,
            "subject": str(subject_name) if subject_name is not None else None,
            "teacher": str(teacher_name) if teacher_name is not None else None,
            "room": str(room_name) if room_name is not None else None,
            "start_time": str(start_time) if start_time is not None else None,
            "end_time": str(end_time) if end_time is not None else None,
            "lesson_id": lesson.id if lesson else None,
            "room_id": room.id if room else None,
            "status": "ready" if not issues else "needs_mapping",
            "issues": issues,
        })

    workbook.close()

    ready_count = sum(
        1 for item in result if item["status"] == "ready"
    )

    return {
        "total_rows": len(result),
        "ready_rows": ready_count,
        "needs_mapping": len(result) - ready_count,
        "can_import": len(result) > 0 and ready_count == len(result),
        "rows": result,
    }


def overlaps(start1, end1, start2, end2):
    """Проверяет пересечение двух временных интервалов."""
    return start1 < end2 and start2 < end1


def minutes_between(start_time, end_time):
    """Количество минут между двумя значениями времени."""
    start_minutes = start_time.hour * 60 + start_time.minute
    end_minutes = end_time.hour * 60 + end_time.minute
    return end_minutes - start_minutes


def check_transition(
    db: Session,
    previous_room: Room,
    target_room: Room,
    available_minutes: int,
):
    """Проверяет, успеет ли преподаватель перейти в другую аудиторию."""

    transition = calculate_transition_time(
        db=db,
        previous_room_id=previous_room.id,
        previous_access_type=previous_room.access_type,
        target_room_id=target_room.id,
        target_access_type=target_room.access_type,
    )

    if transition is None:
        raise HTTPException(
            status_code=409,
            detail=(
                f"Не найден маршрут из аудитории "
                f"{previous_room.name} в аудиторию {target_room.name}. "
                "Невозможно проверить время перехода."
            ),
        )

    required_minutes = transition["total_time"]

    if required_minutes > available_minutes:
        raise HTTPException(
            status_code=409,
            detail=(
                f"Недостаточно времени на переход из аудитории "
                f"{previous_room.name} в аудиторию {target_room.name}. "
                f"Необходимо {required_minutes} мин., "
                f"доступно {available_minutes} мин. "
                f"Время перемещения: {transition['travel_time']} мин.; "
                f"операции с ключами: {transition['key_time']} мин."
            ),
        )


@router.get("/")
def get_weekly_schedule(db: Session = Depends(get_db)):
    entries = (
        db.query(WeeklyScheduleEntry)
        .order_by(
            WeeklyScheduleEntry.weekday,
            WeeklyScheduleEntry.start_time,
        )
        .all()
    )

    result = []

    for entry in entries:
        lesson = db.get(Lesson, entry.lesson_id)
        room = db.get(Room, entry.room_id)

        result.append({
            "id": entry.id,
            "lesson_id": entry.lesson_id,
            "weekday": weekday,
            "start_time": start_time.strftime("%H:%M") if start_time else None,
            "end_time": end_time.strftime("%H:%M") if end_time else None,
            "room_id": entry.room_id,
            "room_name": room.name if room else None,
            "teacher_id": lesson.teacher_id if lesson else None,
            "subject_id": lesson.subject_id if lesson else None,
            "group_id": lesson.group_id if lesson else None,
        })

    return result


@router.post("/")
def assign_lesson(
    data: WeeklyScheduleCreate,
    db: Session = Depends(get_db),
):
    lesson = db.get(Lesson, data.lesson_id)
    room = db.get(Room, data.room_id)

    if not lesson:
        raise HTTPException(404, "Занятие не найдено")

    if not room:
        raise HTTPException(404, "Аудитория не найдена")

    if data.start_time >= data.end_time:
        raise HTTPException(
            400,
            "Время окончания должно быть позже времени начала",
        )

    existing = (
        db.query(WeeklyScheduleEntry)
        .filter(
            WeeklyScheduleEntry.lesson_id == data.lesson_id,
            WeeklyScheduleEntry.weekday == data.weekday,
        )
        .first()
    )

    if existing:
        raise HTTPException(
            status_code=409,
            detail="Это занятие уже добавлено в расписание на этот день."
        )


    entries = (
        db.query(WeeklyScheduleEntry)
        .filter(
            WeeklyScheduleEntry.weekday == data.weekday
        )
        .all()
    )

    # Сначала проверяем конфликты расписания.
    for entry in entries:
        other_lesson = db.get(Lesson, entry.lesson_id)

        if not other_lesson:
            continue

        if not overlaps(
            data.start_time,
            data.end_time,
            entry.start_time,
            entry.end_time,
        ):
            continue

        if entry.room_id == data.room_id:
            raise HTTPException(
                409,
                "Аудитория уже занята в это время",
            )

        if other_lesson.teacher_id == lesson.teacher_id:
            raise HTTPException(
                409,
                "У преподавателя уже есть занятие в это время",
            )

        if other_lesson.group_id == lesson.group_id:
            raise HTTPException(
                409,
                "У группы уже есть занятие в это время",
            )

    # Ищем ближайшее предыдущее и следующее занятие
    # этого же преподавателя в этот день.
    teacher_entries = []

    for entry in entries:
        other_lesson = db.get(Lesson, entry.lesson_id)

        if (
            other_lesson
            and other_lesson.teacher_id == lesson.teacher_id
        ):
            teacher_entries.append((entry, other_lesson))

    previous = [
        (entry, other_lesson)
        for entry, other_lesson in teacher_entries
        if entry.end_time <= data.start_time
    ]

    following = [
        (entry, other_lesson)
        for entry, other_lesson in teacher_entries
        if entry.start_time >= data.end_time
    ]

    # Проверяем переход от предыдущего занятия к новому.
    if previous:
        previous_entry, _ = max(
            previous,
            key=lambda item: item[0].end_time,
        )

        previous_room = db.get(Room, previous_entry.room_id)

        if not previous_room:
            raise HTTPException(
                409,
                "Не удалось найти аудиторию предыдущего занятия",
            )

        available_minutes = minutes_between(
            previous_entry.end_time,
            data.start_time,
        )

        check_transition(
            db=db,
            previous_room=previous_room,
            target_room=room,
            available_minutes=available_minutes,
        )

    # Проверяем переход от нового занятия к следующему.
    if following:
        next_entry, _ = min(
            following,
            key=lambda item: item[0].start_time,
        )

        next_room = db.get(Room, next_entry.room_id)

        if not next_room:
            raise HTTPException(
                409,
                "Не удалось найти аудиторию следующего занятия",
            )

        available_minutes = minutes_between(
            data.end_time,
            next_entry.start_time,
        )

        check_transition(
            db=db,
            previous_room=room,
            target_room=next_room,
            available_minutes=available_minutes,
        )

    # Сохраняем запись только после успешных проверок.
    entry = WeeklyScheduleEntry(
        lesson_id=data.lesson_id,
        weekday=data.weekday,
        room_id=data.room_id,
        start_time=data.start_time,
        end_time=data.end_time,
    )

    db.add(entry)
    db.commit()
    db.refresh(entry)

    return {
        "status": "success",
        "message": "Занятие добавлено в недельное расписание",
        "id": entry.id,
    }


@router.delete("/{entry_id}")
def delete_weekly_entry(
    entry_id: int,
    db: Session = Depends(get_db),
):
    entry = db.get(WeeklyScheduleEntry, entry_id)

    if not entry:
        raise HTTPException(404, "Запись не найдена")

    db.delete(entry)
    db.commit()

    return {"status": "success"}
